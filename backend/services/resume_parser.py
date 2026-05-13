"""
services/resume_parser.py — AI Resume Parser (Phase 3)
=======================================================
Parses PDF and DOCX resumes into structured JSON using GPT-4o.
Falls back to regex extraction if OpenAI is not configured.
"""

import json
import re
from typing import Optional
from loguru import logger


_PARSE_SYSTEM_PROMPT = """
You are an expert resume parser. Extract structured information from the raw resume text.
Return ONLY valid JSON — no markdown, no explanation, no code fences.

JSON schema:
{
  "name": "full name",
  "email": "email",
  "phone": "phone number",
  "location": "city, state",
  "summary": "professional summary",
  "skills": ["skill1", "skill2"],
  "experience": [
    {
      "title": "job title",
      "company": "company name",
      "start_date": "YYYY-MM",
      "end_date": "YYYY-MM or Present",
      "description": "responsibilities",
      "years": 1.5
    }
  ],
  "education": [
    {
      "degree": "degree name",
      "institution": "school name",
      "graduation_year": 2022
    }
  ],
  "total_years_experience": 4.5,
  "linkedin_url": "",
  "portfolio_url": ""
}

Rules:
- Missing fields: use "" for strings, [] for arrays, 0 for numbers.
- Calculate 'years' per job and sum for 'total_years_experience'.
- Extract ALL skills mentioned anywhere.
- Do not invent information not in the text.
"""


class ResumeParser:
    """Parse PDF/DOCX resumes into structured JSON via GPT-4o."""

    def __init__(self):
        self._client = None

    async def parse(self, file_path: str, file_type: str) -> tuple[str, dict]:
        """
        Main entry point. Returns (raw_text, parsed_json).
        """
        logger.info(f"[ResumeParser] Parsing {file_type.upper()}: {file_path}")
        raw_text = self._extract_text(file_path, file_type)
        if not raw_text:
            logger.warning("[ResumeParser] No text extracted")
            return "", self._empty()
        logger.info(f"[ResumeParser] Extracted {len(raw_text)} chars")
        parsed = await self._call_gpt(raw_text)
        logger.success(
            f"[ResumeParser] Done — name='{parsed.get('name')}' "
            f"skills={len(parsed.get('skills', []))} "
            f"exp={len(parsed.get('experience', []))} roles"
        )
        return raw_text, parsed

    # ── Text Extraction ────────────────────────────────────────────────────────

    def _extract_text(self, file_path: str, file_type: str) -> str:
        if file_type == "pdf":
            return self._extract_pdf(file_path)
        elif file_type == "docx":
            return self._extract_docx(file_path)
        return ""

    def _extract_pdf(self, file_path: str) -> str:
        try:
            import fitz
            doc = fitz.open(file_path)
            text = "\n".join(page.get_text() for page in doc)
            doc.close()
            return text
        except ImportError:
            logger.error("[ResumeParser] PyMuPDF not installed — pip install PyMuPDF")
            return ""
        except Exception as e:
            logger.error(f"[ResumeParser] PDF error: {e}")
            return ""

    def _extract_docx(self, file_path: str) -> str:
        try:
            from docx import Document
            doc = Document(file_path)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for tbl in doc.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text.strip())
            return "\n".join(parts)
        except ImportError:
            logger.error("[ResumeParser] python-docx not installed")
            return ""
        except Exception as e:
            logger.error(f"[ResumeParser] DOCX error: {e}")
            return ""

    # ── GPT-4o Parsing ────────────────────────────────────────────────────────

    async def _call_gpt(self, raw_text: str) -> dict:
        client = self._get_client()
        if client is None:
            logger.warning("[ResumeParser] OpenAI not configured — using basic extraction")
            return self._basic_extract(raw_text)
        try:
            truncated = raw_text[:12000]
            response = await client.chat.completions.create(
                model=self._model(),
                messages=[
                    {"role": "system", "content": _PARSE_SYSTEM_PROMPT},
                    {"role": "user", "content": f"Parse this resume:\n\n{truncated}"},
                ],
                temperature=0.1,
                max_tokens=2000,
                response_format={"type": "json_object"},
            )
            data = json.loads(response.choices[0].message.content)
            return self._validate(data)
        except Exception as e:
            logger.error(f"[ResumeParser] GPT call failed: {e} — using fallback")
            return self._basic_extract(raw_text)

    def _get_client(self):
        """Get shared OpenAI/OpenRouter client."""
        if self._client:
            return self._client
        from backend.services.openai_client import get_client
        self._client = get_client()
        return self._client

    def _model(self) -> str:
        from backend.services.openai_client import get_model
        return get_model()

    # ── Validation & Fallback ─────────────────────────────────────────────────

    def _validate(self, data: dict) -> dict:
        base = self._empty()
        for key, default in base.items():
            if key not in data:
                data[key] = default
        if not data.get("total_years_experience") and data.get("experience"):
            data["total_years_experience"] = round(
                sum(e.get("years", 0) for e in data["experience"]), 1
            )
        return data

    def _basic_extract(self, raw_text: str) -> dict:
        """Regex fallback when GPT is unavailable."""
        lines = [l.strip() for l in raw_text.split("\n") if l.strip()]
        result = self._empty()
        if lines:
            result["name"] = lines[0]
        m = re.search(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}", raw_text)
        if m:
            result["email"] = m.group()
        m = re.search(r"(\+?\d{1,3}[\s\-]?)?\(?\d{3}\)?[\s\-]?\d{3}[\s\-]?\d{4}", raw_text)
        if m:
            result["phone"] = m.group().strip()
        m = re.search(r"linkedin\.com/in/[\w\-]+", raw_text, re.IGNORECASE)
        if m:
            result["linkedin_url"] = "https://" + m.group()
        m = re.search(r"github\.com/[\w\-]+", raw_text, re.IGNORECASE)
        if m:
            result["portfolio_url"] = "https://" + m.group()
        result["summary"] = "Extracted via regex fallback (GPT-4o unavailable)"
        return result

    def _empty(self) -> dict:
        return {
            "name": "", "email": "", "phone": "", "location": "", "summary": "",
            "skills": [], "experience": [], "education": [],
            "total_years_experience": 0.0, "linkedin_url": "", "portfolio_url": "",
        }


if __name__ == "__main__":
    import asyncio, sys, os
    TEST = os.environ.get("TEST_RESUME_PATH", "./uploads/test_resume.pdf")
    FTYPE = "pdf" if TEST.endswith(".pdf") else "docx"

    async def main():
        if not os.path.exists(TEST):
            print(f"No resume at {TEST}. Set TEST_RESUME_PATH env var.")
            sys.exit(1)
        parser = ResumeParser()
        raw, parsed = await parser.parse(TEST, FTYPE)
        print(json.dumps(parsed, indent=2))
        checks = [
            ("Name", bool(parsed.get("name"))),
            ("Email", bool(parsed.get("email"))),
            ("Skills", len(parsed.get("skills", [])) > 0),
        ]
        for label, ok in checks:
            print(f"{'✅' if ok else '❌'} {label}")

    asyncio.run(main())
